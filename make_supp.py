import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
m=pd.read_csv('modeling_dataset.csv'); e=pd.read_csv('real_ext/extended_panel.csv'); n=pd.read_csv('real_ext/new_companies.csv')
g=m.groupby('company').agg(src=('source_url','first'),pg=('page_verified','first'),pr=('page_refs','first')).reset_index()
rows=[]
for _,r in e.iterrows():
    if r.origin=='new':
        x=n[n.company==r.company].iloc[0]
        rows.append([r.company,r.macro,x.s1_cur,x.s1_prev,x.s2_cur,x.s2_prev,r.tot_cur,r.tot_prev,'extended (17)',x.source_file,str(x.page),x.note if isinstance(x.note,str) else ''])
    else:
        x=g[g.company==r.company].iloc[0]; pg=x.pg if isinstance(x.pg,str) else (x.pr if isinstance(x.pr,str) else '')
        o='pilot (38)'; note=''
        if r.company=='L&T Ltd.': note='FY2023-24 values taken from the filing file labelled PAM_ (p.36): Scope 1 635,646 / 693,115; Scope 2 349,682 / 273,719'; pg='p.36'
        rows.append([r.company,r.macro,None,None,None,None,r.tot_cur,r.tot_prev,o,x.src,pg,note])
cols=['Company','Macro-sector','Scope 1 current (tCO2e)','Scope 1 previous','Scope 2 current','Scope 2 previous','Scope 1+2 current (tCO2e)','Scope 1+2 previous','Panel','Source file / URL','Location in filing (as recorded)','Note']
wb=Workbook(); ws=wb.active; ws.title='S1 panel'
ws.append(cols)
for r in rows: ws.append([None if (isinstance(v,float) and np.isnan(v)) else v for v in r])
nr=len(rows)+1
ws.cell(1,13,'log10 level'); ws.cell(1,14,'YoY log change')
for i in range(2,nr+1):
    ws.cell(i,13,f'=IF(AND(ISNUMBER(G{i}),G{i}>0),LOG10(G{i}),"")')
    ws.cell(i,14,f'=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i}),G{i}>0,H{i}>0),LN(G{i}/H{i}),"")')
    for c in (3,4,5,6,7,8): ws.cell(i,c).number_format='#,##0.##'
    for c in (13,14): ws.cell(i,c).number_format='0.000'
hf=PatternFill('solid',fgColor='DDDDDD')
for c in range(1,15):
    cell=ws.cell(1,c); cell.font=Font(name='Arial',bold=True,size=10); cell.fill=hf; cell.alignment=Alignment(wrap_text=True,vertical='top')
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.font=Font(name='Arial',size=10)
for i,w in enumerate([34,24,14,14,14,14,16,16,14,48,44,56,11,11],1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes='B2'
n2=wb.create_sheet('Notes')
for t in ['Supplementary Table S1 for: Peer-Conditioned Cross-Modal Consistency Scoring for ESG Disclosure Verification (Krishna and Ahamad).',
 'Scope 1+2 totals are as printed in the BRSR Principle 6 emissions table of each filing. Blank cells mean the filing gave no value; nothing was estimated.',
 'Pilot (38) rows hold totals only (from the audited claim-level dataset); the location column is copied from that dataset and is a page number for some companies and a section name for others.',
 'Extended (17) rows give the source PDF name and printed page index from text extraction (pdftotext page order, which may differ from the printed page number).',
 'Two filings could not be used (no readable emissions table; figures reported as not applicable). One company has no previous-year Scope 2 value.',
 'Extraction caveats (unit plausibility, missing prior-year values, file-name mismatches) are stated in the Note column; values are kept as extracted and were not corrected.',
 'Several local file names in the data folder do not match the company inside (PAM_ = L&T, kavinavora_ = Reliance, SAKSHIKAPUR_ = Tata Steel, Thanos_ = SBI Life).',
 'Columns M and N are formulas: log10 of the current total and ln(current/previous), as used in Section 8 of the paper.']:
    n2.append([t])
n2.column_dimensions['A'].width=150
wb.save('Supplementary_Table_S1_BRSR_Panel.xlsx')
print(nr-1,'rows')
